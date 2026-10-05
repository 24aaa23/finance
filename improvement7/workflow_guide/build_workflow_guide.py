"""Build the guide and editable diagrams without changing pipeline code."""
from pathlib import Path
from collections import Counter
import json
import math
import xml.etree.ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent
QA = OUT / 'qa'
QA.mkdir(exist_ok=True)
FONT = Path('C:/Windows/Fonts/arial.ttf')
BASELINE = {
    'benchmark': 'timing_20261003_134611_19444.json',
    'at_risk_critical': 'timing_20261003_134616_24944.json',
    'batch_status': 'timing_20261003_134623_27916.json',
    'enrichment_context': 'timing_20261003_134628_460.json',
    'multi_step_comparative': 'timing_20261003_134632_12608.json',
    'reference_compliance': 'timing_20261003_134638_17912.json',
    'scoring_quantitative': 'timing_20261003_134643_10128.json',
    'temporal_transaction': 'timing_20261003_134648_28192.json',
}
metrics = []
stages = Counter()
for kind, filename in BASELINE.items():
    log = json.loads((ROOT / '.runtime/logs' / filename).read_text(encoding='utf-8'))
    assert log['status'] == 'completed' and len(log['questions']) == 125
    calls = sum(q['llm_calls'] for q in log['questions'])
    seconds = sum(q['processing_seconds'] for q in log['questions'])
    for call in log['llm_calls']:
        if call.get('query_id'):
            stages[call['stage']] += 1
    metrics.append(dict(type=kind, questions=125, calls=calls,
                        calls_per_question=calls / 125, seconds_per_question=seconds / 125,
                        source=filename))
total_calls = sum(r['calls'] for r in metrics)
(OUT / 'workflow_metrics.json').write_text(json.dumps({
    'snapshot_date': '2026-10-03', 'cohort': 'Original eight completed full runs at 13:46',
    'exclusions': 'Startup calls, later reruns, interrupted runs and grading',
    'questions': 1000, 'calls': total_calls, 'calls_per_question': total_calls / 1000,
    'by_type': metrics, 'stages': dict(stages)}, indent=2), encoding='utf-8')

COLORS = {'llm': '#E3EFF9', 'local': '#E5F2EA', 'data': '#F1F2F4', 'repair': '#FFF1CE'}
diagrams = []

def diagram(name, width, height, nodes, edges, labels):
    im = Image.new('RGB', (width, height), 'white')
    d = ImageDraw.Draw(im)
    by_id = {n[0]: n for n in nodes}
    # The image and draw.io source share geometry, so edits have a clear starting point.
    for start, end, route in edges:
        a, b = by_id[start], by_id[end]
        if route == 'down':
            pts = [(a[1] + a[3] / 2, a[2] + a[4]), (b[1] + b[3] / 2, b[2])]
        elif route == 'wrap':
            y = (a[2] + a[4] + b[2]) / 2
            pts = [(a[1] + a[3] / 2, a[2] + a[4]), (a[1] + a[3] / 2, y),
                   (b[1] + b[3] / 2, y), (b[1] + b[3] / 2, b[2])]
        else:
            pts = [(a[1] + a[3], a[2] + a[4] / 2), (b[1], b[2] + b[4] / 2)]
        d.line(pts, fill='#55616B', width=4)
        x, y = pts[-1]
        px, py = pts[-2]
        angle = math.atan2(y - py, x - px)
        d.polygon([(x, y), (x - 15 * math.cos(angle - .5), y - 15 * math.sin(angle - .5)),
                   (x - 15 * math.cos(angle + .5), y - 15 * math.sin(angle + .5))], fill='#55616B')
    for ident, x, y, w, h, txt, category in nodes:
        d.rounded_rectangle((x, y, x + w, y + h), radius=8, fill=COLORS[category], outline='#71808D', width=2)
        font = ImageFont.truetype(str(FONT), 30)
        bounds = d.multiline_textbbox((0, 0), txt, font=font, spacing=8, align='center')
        assert bounds[2] < w - 12 and bounds[3] - bounds[1] < h - 12, (ident, bounds, w, h)
        d.multiline_text((x + w / 2, y + h / 2), txt, font=font, fill='#172126',
                         anchor='mm', spacing=8, align='center')
    for x, y, txt in labels:
        d.text((x, y), txt, font=ImageFont.truetype(str(FONT), 29), fill='#172126')
    path = QA / (name + '.png')
    im.save(path)
    diagrams.append((name, width, height, nodes, edges, labels))
    return path

main_nodes = [
    ('docs', 10, 45, 420, 115, 'Domain docs + YAML\nPhysical SQLite schema', 'data'),
    ('knowledge', 490, 45, 420, 115, 'Domain Knowledge Agent\nCache 0; cold cache 1-3', 'llm'),
    ('context', 970, 45, 420, 115, 'Compiled business context\nReused across questions', 'data'),
    ('qu', 10, 300, 420, 115, 'Query Understanding\nQuestion -> contract | 1 call', 'llm'),
    ('retrieve', 490, 300, 420, 115, 'Retrieve\nSelect source tables | 1 call', 'llm'),
    ('decompose', 970, 300, 420, 115, 'Decompose\nPlan source branches | 1 call', 'llm'),
    ('qs', 10, 555, 246, 140, 'Query_Spec\nRaw fields\n1 call / branch', 'llm'),
    ('generate', 294, 555, 246, 140, 'Generate\nLocal 0 calls\nLLM fallback 1', 'llm'),
    ('pre', 578, 555, 246, 140, 'Pre_Scan\nValidate SQL\n0 calls', 'local'),
    ('scan', 862, 555, 246, 140, 'Scan\nRead SQLite\n0 calls', 'local'),
    ('process', 1146, 555, 244, 140, 'Processing\nPreserve rows\n0 calls', 'local'),
    ('final', 10, 835, 420, 115, 'Final_Spec\nCalculation plan | 1 call', 'llm'),
    ('execute', 490, 835, 420, 115, 'Compile and execute\nJoins + calculations | 0 calls', 'local'),
    ('validate', 970, 835, 420, 115, 'Validate + Explain\n0 default; up to 1 each', 'llm'),
    ('report', 10, 1080, 660, 105, 'Save result or error to raw report\nThen continue with the next question', 'data'),
    ('grader', 750, 1080, 640, 105, 'Separate grader\nCompare saved answer with ground truth', 'data'),
]
main_edges = [('docs', 'knowledge', 'right'), ('knowledge', 'context', 'right'),
              ('context', 'qu', 'wrap'), ('qu', 'retrieve', 'right'), ('retrieve', 'decompose', 'right'),
              ('decompose', 'qs', 'wrap'), ('qs', 'generate', 'right'), ('generate', 'pre', 'right'),
              ('pre', 'scan', 'right'), ('scan', 'process', 'right'), ('process', 'final', 'wrap'),
              ('final', 'execute', 'right'), ('execute', 'validate', 'right'),
              ('validate', 'report', 'wrap'), ('report', 'grader', 'right')]
main_png = diagram('Main_workflow', 1400, 1220, main_nodes, main_edges, [
    (10, 0, 'STARTUP   Read-only database checks and shared business context'),
    (10, 180, 'EACH QUESTION   Original question + shared context enter planning'),
    (10, 435, 'REPEAT THIS ROW FOR EACH SOURCE BRANCH   B branches'),
    (10, 715, 'ALL BRANCHES COME TOGETHER   Deterministic calculation runtime'),
    (10, 970, 'OUTPUT   Raw report first; grading is a separate operation'),
])
retry_png = diagram('Query_understanding_recovery', 1400, 625, [
    ('attempt', 10, 55, 410, 115, 'Attempt 1\nInterpret the question', 'llm'),
    ('check', 500, 55, 400, 115, 'Local checks + cleanup\nIs the contract valid?', 'local'),
    ('ok', 980, 55, 410, 115, 'YES -> use contract\nContinue to Retrieve', 'local'),
    ('retry', 500, 290, 400, 130, 'NO -> feedback + retry\nAttempts 2 and 3\nOnly when needed', 'repair'),
    ('valid', 980, 290, 410, 130, 'Valid on retry\nUse contract and continue', 'local'),
    ('fallback', 500, 510, 890, 110, 'Still invalid after attempt 3 -> return empty contract\nContinue base pipeline using the original question', 'repair'),
], [('attempt', 'check', 'right'), ('check', 'ok', 'right'), ('check', 'retry', 'down'),
    ('retry', 'valid', 'right'), ('retry', 'fallback', 'down')], [])

repair_nodes = []
repair_edges = []
repair_lanes = [
    ('Pre-scan check fails', 'Generate again\nLocal 0 or LLM 1', 'Pre-scan check again\nThen Scan if valid'),
    ('Scan fails\nUnknown source fields', 'Query_Spec again\n1 LLM call / invocation', 'Generate -> pre-scan\nThen Scan again'),
    ('Scan fails\nOther SQL error', 'Refine\n1 LLM call / invocation', 'Pre-scan -> Scan again\nMay also regenerate SQL'),
    ('Final plan fails\nCompile or execution error', 'Final_Spec again\n1 LLM call / invocation', 'Compile + execute again\nUp to 4 plan attempts'),
    ('Missing upstream data\nRequires retrieval repair', 'Decompose again\n1 LLM call / invocation', 'Rerun source branches\nThen Final_Spec again'),
    ('Optional consultation\nRequested by a plan stage', 'Domain advice 1 call\nOr understanding 1-3', 'Reask requesting stage\nOr restart if meaning changes'),
]
for i, texts in enumerate(repair_lanes):
    y = 15 + i * 190
    for j, (x, w, text_value) in enumerate(zip([10, 490, 970], [420, 420, 420], texts)):
        repair_nodes.append((f'r{i}_{j}', x, y, w, 125, text_value, 'llm' if j == 1 else ('repair' if j == 0 else 'local')))
    repair_edges.extend([(f'r{i}_0', f'r{i}_1', 'right'), (f'r{i}_1', f'r{i}_2', 'right')])
repair_png = diagram('All_self_healing_paths', 1400, 1105, repair_nodes, repair_edges, [])

mxfile = ET.Element('mxfile', host='app.diagrams.net', type='device')
for name, width, height, nodes, edges, labels in diagrams:
    page = ET.SubElement(mxfile, 'diagram', name=name)
    model = ET.SubElement(page, 'mxGraphModel', page='1', pageWidth=str(width), pageHeight=str(height))
    root = ET.SubElement(model, 'root')
    ET.SubElement(root, 'mxCell', id='0')
    ET.SubElement(root, 'mxCell', id='1', parent='0')
    for ident, x, y, w, h, txt, category in nodes:
        cell = ET.SubElement(root, 'mxCell', id=ident, value=txt, vertex='1', parent='1',
            style=f'rounded=1;whiteSpace=wrap;html=0;arcSize=5;fontSize=30;fontColor=#172126;fillColor={COLORS[category]};strokeColor=#71808D;spacing=8;')
        ET.SubElement(cell, 'mxGeometry', x=str(x), y=str(y), width=str(w), height=str(h), **{'as': 'geometry'})
    for i, (a, b, route) in enumerate(edges):
        direction = 'exitX=0.5;exitY=1;entryX=0.5;entryY=0;' if route != 'right' else 'exitX=1;exitY=0.5;entryX=0;entryY=0.5;'
        cell = ET.SubElement(root, 'mxCell', id=f'edge{i}', edge='1', parent='1', source=a, target=b,
                             style='edgeStyle=orthogonalEdgeStyle;rounded=0;endArrow=block;strokeColor=#55616B;strokeWidth=3;' + direction)
        ET.SubElement(cell, 'mxGeometry', relative='1', **{'as': 'geometry'})
    for i, (x, y, txt) in enumerate(labels):
        cell = ET.SubElement(root, 'mxCell', id=f'label{i}', value=txt, vertex='1', parent='1',
                             style='text;html=0;align=left;fontSize=29;fontColor=#172126;')
        ET.SubElement(cell, 'mxGeometry', x=str(x), y=str(y), width='1380', height='38', **{'as': 'geometry'})
ET.indent(mxfile)
ET.ElementTree(mxfile).write(OUT / 'Improvement7_Workflow.drawio', encoding='utf-8', xml_declaration=True)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.05
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
header = sec.header.paragraphs[0]
header.text = 'IMPROVEMENT 7  |  WORKFLOW GUIDE'
header.runs[0].font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Code and historical logs inspected 3 October 2026  |  ' ).font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def heading(text):
    doc.add_heading(text, level=2)

def page(title):
    para = doc.add_heading(title, level=1)
    para.paragraph_format.page_break_before = True
    para.paragraph_format.space_before = Pt(0)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, values in enumerate([headers] + list(rows)):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '37474F' if i == 0 else ('F2F5F6' if i % 2 == 0 else 'FFFFFF'))
            props.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), '95')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            props.append(margins)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.0
            if j > 0 and widths[j] < 1.2:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(str(value))
            r.font.size = Pt(10.5)
            if i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
        if i == 0:
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + side)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + key), value)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

doc.add_heading('Improvement 7 Workflow Guide', 0)
p('Follow a question from interpretation to a saved answer. This guide explains the agents, operators, recovery limits, measured model calls and the files that control the workflow.')
doc.add_picture(str(main_png), width=Inches(7.1))
p('Blue = stage that can call a model, not a guarantee that it does. Green = local code. Gray = inputs or outputs. This is the forward path; the separate recovery map shows repeated calls and consultation. Generate uses an LLM only if local rendering cannot produce SQL.')
p('Read left to right within each row, then follow the arrow to the next row. Branches are logical data paths; the executor does not automatically run them in parallel.', bold=True)

page('What the two agents do')
heading('Domain Knowledge Agent learns the shared rules')
p('At startup, this agent reads the authorized domain documents and schema. It produces a reusable, cited business knowledge pack. A matching cached pack avoids another model call. It does not read the test answer or query database rows.')
table(['Input', 'What it contributes'], [
    ('domain_intro_latest.prompt', 'General wealth management concepts and domain context.'),
    ('business_rules_addendum.md', 'Business definitions, formulas and rules to follow.'),
    ('table_medatada/*.yaml', 'Table and column meanings, synonyms and enum descriptions.'),
    ('SQLite physical schema', 'Which tables and columns actually exist.')], [2.85, 4.25])
p('The YAML files describe the data; they are not an extra agent and do not execute rules by themselves. The physical database remains the authority for real table and column names.')
heading('Query Understanding Agent interprets one question')
p('For each question, it uses the annotated schema and business knowledge to propose a contract: which fields mean what, needed metrics, filters, grouping, projected columns, ordering and limits. Later stages can use this contract to stay aligned with the question.')
heading('A simple illustrative question')
p('Example: "Show total transaction amount per client for last month." This is illustrative, not a benchmark answer or a promise about exact column names.')
p('1. Domain context supplies definitions such as what a client, amount and transaction date mean. Query Understanding binds those meanings to actual schema fields and identifies the time filter and per-client grouping.')
p('2. Retrieve selects the relevant tables. Decompose creates the source branches. Query_Spec requests raw IDs, amounts and dates; Generate and Scan fetch the rows.')
p('3. Final_Spec plans any needed join, filter and sum. Local operators execute that plan. Validate and Explain prepare the result for the report.')
heading('What happens when another stage asks an agent')
p('Most stages receive the already compiled context; receiving context is not another agent call. Optional live consultation is different: with agent consultation enabled, Decompose, Query_Spec and Final_Spec can request focused advice. Generate and Scan do not directly request that advice. See the recovery page for the budget and restart behavior.')

page('Important operators')
p('An operator is one named step in the pipeline. A model-assisted operator proposes a plan; local Python and SQLite perform the actual data work.')
table(['Operator', 'Job', 'Normal model calls'], [
    ('Retrieve', 'Select source tables from schema.', '1'),
    ('Decompose', 'Split the task into source branches.', '1'),
    ('Query_Spec', 'Choose raw fields and retrieval filters.', '1 / branch'),
    ('Generate', 'Render raw SQL locally, or use 1 LLM call if rendering returns None.', '0 or 1'),
    ('Pre_Scan_Validate', 'Check SQL, schema and raw retrieval rules.', '0'),
    ('Scan', 'Execute read-only SQLite retrieval.', '0'),
    ('Refine', 'Not on the normal path. Each invocation makes 1 LLM call.', 'Repair: 1'),
    ('Processing_Spec', 'Keep branch rows and schema available.', '0'),
    ('Final_Spec', 'Build the final calculation plan as JSON.', '1'),
    ('Validate / Explain', 'Check and format the executed answer.', '0 by default'),
], [1.6, 4.3, 1.2])
heading('Calculation operators inside the final plan')
table(['Operator family', 'Meaning'], [
    ('Integrate', 'Join source datasets and check join cardinality.'),
    ('Filter_Aggregate', 'Filter and aggregate using sum, average, count or distinct count.'),
    ('Math_Compute', 'Calculate derived values and arithmetic expressions.'),
    ('Date_Extract / Bucket', 'Extract date parts or assign values to defined buckets.'),
    ('Order_By / Distinct', 'Sort or limit output; remove duplicates.'),
    ('Set_Intersect / Difference / Union', 'Combine or compare sets of rows.')], [2.65, 4.45])
p('These calculation operators make zero model calls in the current compiled data path. Integrate, Filter_Aggregate and Order_By also retain legacy LLM fallbacks outside that path. A successful execution can still be a wrong answer; the separate grader checks correctness.')

page('Self healing and its limits')
doc.add_picture(str(retry_png), width=Inches(7.1))
p('Query Understanding stops retrying as soon as a valid contract is available. After three invalid attempts, current code uses the base pipeline with the original question; it does not select an invalid "best" contract.')
table(['Stage', 'Bound and behavior'], [
    ('Domain compilation', 'Cache miss: up to 3 validation attempts. If all fail, startup stops.'),
    ('Query_Spec', 'Up to 3 attempts per branch in a decomposition pass.'),
    ('Pre-scan and Scan', 'Default 3 checks / attempts per loop, with SQL regeneration or Refine when applicable.'),
    ('Final_Spec', 'Up to 4 attempts per execution; structured errors guide repair.'),
    ('Decomposition', 'Up to 3 passes; missing source data or upstream planning errors can restart branches.')], [1.75, 5.35])
heading('Optional specialist consultation')
p('Default off. With --agent-consultation on, Decompose, Query_Spec and Final_Spec may request domain or interpretation advice. The shared budget is 5 specialist requests per question. A domain request uses one model call; a Query Understanding request can use up to three attempts. A changed interpretation restarts retrieval and decomposition within the existing pass budget.')
heading('Why a question or batch can still stop')
p('Retry limits apply to individual loops, not to the entire question. Nested retries can add many calls. Exhausted SQL or calculation repairs produce an error row and the batch normally moves on. An API quota failure can stop the batch; invalid startup knowledge, database health or report-lock problems can prevent it from starting. Transport exceptions are not guaranteed to reach the Query Understanding fallback.')
p('Best effort means continuing where a supported fallback exists, not inventing database results or silently discarding required constraints. Valid empty results can be successful. Zero execution errors cannot be guaranteed.', bold=True)

page('All self healing paths')
p('Each row below is a conditional recovery route, not six steps that always run. Failed attempts can add both a repair call and repeated downstream calls. Query Understanding has its own recovery diagram on the preceding page.')
doc.add_picture(str(repair_png), width=Inches(7.1))
p('Pre-scan and Scan default to 3 checks or attempts per loop. Decomposition has up to 3 passes. These loops can nest, so the whole question is not limited to 3 model calls. Successful checks stop their loop early.')
p('Generate always tries render_raw_sql first, including during regeneration. A self-healing event is therefore not automatically a model call. Refine, however, calls the model whenever that repair operator is invoked.', bold=True)
p('When a timeout rewrite is too similar to the failed SQL, the executor calls Generate as well. If an unknown-field Query_Spec repair is invalid, it can fall through to Refine. Missing SQL triggers Generate, pre-scan checks and another Scan.')
p('Consultation can add the requesting operator call, the specialist call(s), and another operator call. A revised interpretation instead restarts Retrieve and Decompose. The 5-request budget is not a 5-model-call budget.')

page('Complete model call inventory')
p('A model-call site is not the same as an operator invocation. This inventory separates normal calls, conditional calls and compatibility code not used by the current compiled execution path.')
table(['Call site', 'When the model is called'], [
    ('Domain compilation', 'Cache miss: 1 call per attempt, up to 3. Cache hit: 0.'),
    ('Query Understanding', '1 call per attempt, up to 3 per invocation. Consultation can invoke it again.'),
    ('Retrieve', '1 per invocation; called again when specialist advice changes interpretation.'),
    ('Decompose', '1 per invocation; outer repair passes and consultation can repeat it.'),
    ('Query_Spec', '1 per invocation. Up to 3 branch validation attempts, plus scan repairs, restarts and consultation.'),
    ('Generate', '0 if render_raw_sql succeeds; otherwise 1. The same condition applies on regeneration.'),
    ('Refine', '1 each time the SQL repair operator is called.'),
    ('Final_Spec', '1 per invocation. Up to 4 plan attempts per execution, plus restarts and consultation.'),
    ('Domain consultation', '1 per specialist request when consultation is enabled.'),
    ('Validate', '1 if semantic review is enabled and local execution checks pass; otherwise 0.'),
    ('Explain', '1 when neither the empty/context bypass nor structured JSON path applies.'),
    ('Legacy calculation fallbacks', 'Integrate, Filter_Aggregate and Order_By contain model fallbacks; current compiled data execution bypasses them.')], [1.8, 5.3])
p('Current compiled execution supplies strict_spec=True and actual datasets. Integrate returns via local relational code; Filter_Aggregate and Order_By use their data branches. Their model fallbacks are real source-code sites, but are not additional calls in this workflow.')
p('performance.py wraps and records SDK calls; it is not an extra independent reasoning step. SDK-internal transport retries may make more network attempts than the recorded invocation count. Grading runs separately and adds its own calls.')
p('Source audit: llm_operators/*.py; knowledge.py; query_understanding.py; consultation.py; executor.py; execution.py; main.py; performance.py. The first-page counts describe a successful forward pass, not the full repair budget.')

page('Model calls and measured averages')
heading('Normal path formula')
p('With Query Understanding on, cached domain knowledge, local SQL rendering, consultation off, and default validation and explanation:')
p('Calls per question = 1 Query Understanding + 1 Retrieve + 1 Decompose + B Query_Spec + 1 Final_Spec = B + 4', bold=True)
p('B is the number of source branches, not necessarily the number of unique tables. One branch needs 5 calls; two need 6; four need 8. These are normal-path counts, not hard caps.')
p('Add calls for failed attempts, live consultation, model SQL fallback, optional semantic review or model explanation. A successful cold-cache domain compilation usually adds one shared startup call; it can try up to three times. Turning Query Understanding off removes its normal one call.')
heading('Original eight full runs')
p('Measured from the eight completed timing logs started at 13:46 on 3 October 2026. Every row includes 125 original question attempts, including failed questions. Later reruns and grading are excluded. These are historical measurements, not a new run of the current code.')
table(['Question type', 'Questions', 'Calls', 'Calls / question', 'Seconds / question'], [
    (r['type'].replace('_', ' '), r['questions'], r['calls'], f"{r['calls_per_question']:.2f}", f"{r['seconds_per_question']:.1f}")
    for r in metrics] + [('All types', 1000, total_calls, f'{total_calls / 1000:.2f}',
        f"{sum(r['seconds_per_question'] for r in metrics) / 8:.1f}")], [2.65, .85, .8, .95, 1.85])
p('Across the 1,000 original attempts: 7,263 model calls, averaging 7.26 per question. Average recorded processing time was 57.7 seconds per question; this is not the total elapsed time when multiple terminals run together.', bold=True)
p('A call means one instrumented SDK invocation. Internal SDK retries are not separate entries; their waiting time can be included. Startup calls have no question ID and are excluded from this average. Grader calls are also excluded.')
p('Audit source: improvement7/.runtime/logs/timing_20261003_1346*.json. workflow_metrics.json lists the exact eight files, totals and stage counts so this snapshot can be checked.')

page('How to change the workflow')
heading('Edit the picture or edit the running pipeline')
p('The Word text is editable. Its diagrams are images for stable layout. Open Improvement7_Workflow.drawio in a compatible diagram editor to edit boxes and arrows. Its three pages show the main workflow, Query Understanding recovery and all self-healing paths. Editing these documents does not change pipeline behavior.')
p('For actual behavior changes, use the code map below. Paths are relative to improvement7/run_pipeline_v2_sql unless stated otherwise.')
table(['Change you want', 'File or entry point'], [
    ('Operator order or branches', 'planner.py: AdvancedAOPPlanner._build_three_spec_dag'),
    ('Startup, agent wiring, outer retries and reports', 'main.py: _main and run_single_test'),
    ('Stage dispatch, local retries and repair routing', 'executor.py: AOPExecutor.execute_dag'),
    ('Question contract and its 3-attempt fallback', 'query_understanding.py: understand_query'),
    ('Domain compilation, cache and citations', 'knowledge.py: compile_knowledge'),
    ('Live advice and its shared request budget', 'consultation.py: ConsultationSession'),
    ('Final plan prompting and supported calculations', 'llm_operators/final_spec.py; spec_runtime.py; spec_contracts.py; non_llm_operators/'),
    ('Defaults and command-line switches', 'common.py and run.py'),
    ('Business meaning without changing the DAG', '../business_rules_addendum.md; ../domain_intro_latest.prompt; ../table_medatada/*.yaml')], [2.65, 4.45])
heading('Make a small controlled change')
p('1. Decide whether the change is business meaning, interpretation, operator order or execution. Change only the owning layer. Adding a new operator also requires registry wiring, compiler support, prompt updates and tests; an arrow alone is not enough.')
p('2. Keep the test set and ground truth out of agent prompts and business rules. Define general behavior, then test it without encoding individual benchmark answers.')
p('3. Run the existing unit tests from the financial folder:')
code = p('python -B -m unittest discover -s improvement7/run_pipeline_v2_sql/tests -v')
code.runs[0].font.name = 'Consolas'
code.runs[0].font.size = Pt(9)
p('4. Use a separate output report for a full comparison after changing code. In-place error retries keep successful rows from older runs, so that report is a mixed-run result, not a clean full rerun.')
p('This guide was generated from the local code and saved timing logs. No pipeline logic, business rules, test questions or existing results were changed for this documentation task.')

doc.core_properties.title = 'Improvement 7 Workflow Guide'
doc.core_properties.subject = 'Pipeline agents, operators, retries and measured model calls'
doc.core_properties.author = 'Workflow documentation'
doc.save(OUT / 'Improvement7_Workflow_Guide.docx')
from update_main_diagram import update_main_diagram
update_main_diagram()
print(json.dumps({'document': str(OUT / 'Improvement7_Workflow_Guide.docx'),
                  'calls': total_calls, 'average': total_calls / 1000, 'stages': dict(stages)}))
